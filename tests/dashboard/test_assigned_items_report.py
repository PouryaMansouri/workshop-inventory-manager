"""Tests for assigned items report and XLSX export."""
from datetime import timedelta

import pytest
from django.contrib.auth.models import User
from django.test import Client
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from io import BytesIO

from apps.inventory.models import (
    Category,
    ItemCheckConfirmation,
    ItemInstance,
    Location,
    Product,
    UnitOfMeasure,
    WorkshopUnit,
)
from apps.transactions.models import Transaction
from core.selectors.transactions import TransactionSelector


@pytest.mark.django_db
class TestAssignedItemsReport:
    def setup_method(self):
        self.client = Client()

    def _build_fixture(self):
        workshop = WorkshopUnit.objects.create(name='Mechanic', code='MC')
        category = Category.objects.create(name='Hammer', code='HAM')
        uom = UnitOfMeasure.objects.create(name='Piece', symbol='pcs')
        location = Location.objects.create(name='Main')

        product = Product.objects.create(
            name='Hammer 200',
            category=category,
            workshop_unit=workshop,
            unit_of_measure=uom,
            minimum_stock=0,
        )
        item = ItemInstance.objects.create(
            product=product,
            serial_number='001',
            location=location,
            status='checked_out',
        )

        admin = User.objects.create_user(username='adminx', password='x', is_staff=True)
        worker = User.objects.create_user(username='workerx', password='x', first_name='Ali', last_name='Rezaei')

        active_txn = Transaction.objects.create(
            item=item,
            user=worker,
            transaction_type='checkout',
            date_out=timezone.now() - timedelta(days=2),
            created_by=admin,
        )

        ItemCheckConfirmation.objects.create(item=item, confirmed_by=admin, notes='checked')
        return worker, active_txn

    def test_selector_returns_holding_range_and_last_check(self):
        worker, active_txn = self._build_fixture()
        rows = TransactionSelector.get_assigned_items_rows(holder_id=worker.id)

        assert len(rows) == 1
        row = rows[0]
        assert row['holding_range_start'] == active_txn.date_out
        assert row['holding_range_end'] >= row['holding_range_start']
        assert row['last_check_by'] is not None

    def test_report_export_xlsx(self):
        worker, _ = self._build_fixture()
        viewer = User.objects.create_user(username='viewer', password='x', is_staff=True)
        assert self.client.login(username='viewer', password='x')

        response = self.client.get(reverse('assigned_items_report_export_xlsx'), {'holder_id': worker.id})
        assert response.status_code == 200
        assert response['Content-Type'].startswith('application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

        wb = load_workbook(BytesIO(response.content))
        ws = wb.active
        assert ws.cell(1, 1).value == 'نام تحویل گیرنده'
        assert ws.max_row >= 2
