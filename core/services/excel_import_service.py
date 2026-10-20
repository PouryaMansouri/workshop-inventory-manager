"""Excel import service for real_docs bootstrap."""
from __future__ import annotations

from pathlib import Path
from typing import Any

from django.contrib.auth.models import User
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook

from apps.inventory.models import (
    Category,
    ItemCheckConfirmation,
    ItemInstance,
    Location,
    Product,
    UnitOfMeasure,
    WorkshopUnit,
)
from apps.transactions.models import Transaction


class ExcelImportService:
    """Import bootstrap data from real Excel documents."""

    REQUIRED_FILES = [
        "categories-instrument.xlsx",
        "categories-equipment.xlsx",
        "example.xlsx",
        "coding.xlsx",
        "نمونه لیست اقلام تجویلی.xlsx",
    ]

    @staticmethod
    @transaction.atomic
    def import_from_real_docs(base_path: str | Path) -> dict[str, Any]:
        base = Path(base_path)
        result = {
            "success": True,
            "created_count": 0,
            "error_count": 0,
            "errors": [],
        }

        missing = [name for name in ExcelImportService.REQUIRED_FILES if not (base / name).exists()]
        if missing:
            result["success"] = False
            result["error_count"] = len(missing)
            result["errors"] = [f"Missing file: {name}" for name in missing]
            return result

        try:
            ExcelImportService._reset_operational_data()

            default_uom, _ = UnitOfMeasure.objects.get_or_create(name="عدد", defaults={"symbol": "عدد"})
            main_location, _ = Location.objects.get_or_create(name="انبار اصلی")

            categories = ExcelImportService._import_categories(base)
            created_products, created_items = ExcelImportService._import_example(base, categories, default_uom, main_location)
            assignments = ExcelImportService._import_assignments(base, created_products)

            result["created_count"] = {
                "categories": categories,
                "products": created_products,
                "items": created_items,
                "assignments": assignments,
            }
        except Exception as exc:
            result["success"] = False
            result["error_count"] += 1
            result["errors"].append(str(exc))
        return result

    @staticmethod
    def _reset_operational_data() -> None:
        """Clear inventory and transaction data while preserving auth users."""
        Transaction.objects.all().delete()
        ItemCheckConfirmation.objects.all().delete()
        ItemInstance.objects.all().delete()
        Product.objects.all().delete()
        Category.objects.all().delete()
        WorkshopUnit.objects.all().delete()
        Location.objects.all().delete()
        UnitOfMeasure.objects.all().delete()

    @staticmethod
    def _import_categories(base: Path) -> int:
        created = 0
        code_seq = {"INS": 0, "EQP": 0}

        def import_file(file_name: str, prefix: str) -> int:
            local_created = 0
            wb = load_workbook(base / file_name, data_only=True)
            for sheet in wb.sheetnames:
                ws = wb[sheet]
                for row_idx in range(4, ws.max_row + 1):
                    name = ws.cell(row_idx, 2).value
                    if not name:
                        continue
                    category_name = str(name).strip()
                    if not category_name:
                        continue
                    code_seq[prefix] += 1
                    code = f"{prefix}{code_seq[prefix]:03d}"
                    _, is_created = Category.objects.get_or_create(
                        name=category_name,
                        defaults={"code": code},
                    )
                    if is_created:
                        local_created += 1
            return local_created

        created += import_file("categories-instrument.xlsx", "INS")
        created += import_file("categories-equipment.xlsx", "EQP")
        return created

    @staticmethod
    def _ensure_workshop_unit(code: str) -> WorkshopUnit:
        return WorkshopUnit.objects.get_or_create(
            code=code,
            defaults={"name": f"واحد {code}"},
        )[0]

    @staticmethod
    def _category_for_product_name(product_name: str) -> Category:
        for cat in Category.objects.all():
            if cat.name in product_name:
                return cat
        return Category.objects.order_by("id").first()

    @staticmethod
    def _import_example(base: Path, _categories_created: int, default_uom: UnitOfMeasure, location: Location) -> tuple[int, int]:
        wb = load_workbook(base / "example.xlsx", data_only=True)
        ws = wb[wb.sheetnames[0]]

        created_products = 0
        created_items = 0

        for row_idx in range(5, ws.max_row + 1):
            code = ws.cell(row_idx, 18).value
            item_name = ws.cell(row_idx, 19).value
            if not code or not item_name:
                continue

            code_s = str(code).strip()
            name_s = str(item_name).strip()
            if not code_s or not name_s:
                continue

            unit_code = code_s[1:3] if len(code_s) >= 3 else "GEN"
            workshop = ExcelImportService._ensure_workshop_unit(unit_code)
            category = ExcelImportService._category_for_product_name(name_s)
            if category is None:
                category = Category.objects.create(name="عمومی", code="GEN")

            product, was_created = Product.objects.get_or_create(
                sku_code=code_s,
                defaults={
                    "name": name_s,
                    "category": category,
                    "unit_of_measure": default_uom,
                    "minimum_stock": 0,
                    "is_consumable": False,
                },
            )
            if was_created:
                created_products += 1

            _, item_created = ItemInstance.objects.get_or_create(
                product=product,
                serial_number="001",
                defaults={
                    "location": location,
                    "workshop_unit": workshop,
                    "status": "available",
                },
            )
            if item_created:
                created_items += 1

        return created_products, created_items

    @staticmethod
    def _import_assignments(base: Path, created_products: int) -> int:
        wb = load_workbook(base / "نمونه لیست اقلام تجویلی.xlsx", data_only=True)
        ws = wb[wb.sheetnames[0]]

        responsible_name = str(ws.cell(2, 3).value or "").strip() or "مسئول"
        first_name, last_name = ExcelImportService._split_name(responsible_name)
        username = ExcelImportService._make_username(first_name, last_name)
        holder, _ = User.objects.get_or_create(
            username=username,
            defaults={
                "first_name": first_name,
                "last_name": last_name,
                "is_active": True,
            },
        )

        assigned = 0
        now = timezone.now()
        products = list(Product.objects.select_related("category").all())

        for row_idx in range(6, ws.max_row + 1):
            product_name = ws.cell(row_idx, 3).value
            qty = ws.cell(row_idx, 4).value
            if not product_name or not qty:
                continue

            product_name = str(product_name).strip()
            try:
                qty = int(qty)
            except (TypeError, ValueError):
                continue

            matched = [p for p in products if product_name in p.name or p.name in product_name]
            if not matched:
                continue

            for product in matched[:1]:
                available_items = list(
                    ItemInstance.objects.filter(product=product, status="available").order_by("id")[:qty]
                )
                for item in available_items:
                    item.status = "checked_out"
                    item.save(update_fields=["status"])
                    Transaction.objects.create(
                        item=item,
                        user=holder,
                        transaction_type="checkout",
                        date_out=now,
                        created_by=holder,
                        notes="Bootstrap from real_docs sample",
                    )
                    assigned += 1

        return assigned

    @staticmethod
    def _split_name(full_name: str) -> tuple[str, str]:
        parts = full_name.split()
        if not parts:
            return "کاربر", "نمونه"
        if len(parts) == 1:
            return parts[0], ""
        return parts[0], " ".join(parts[1:])

    @staticmethod
    def _make_username(first_name: str, last_name: str) -> str:
        raw = f"{first_name}.{last_name}".strip(".").replace(" ", "_")
        raw = raw or "assigned_user"
        base = "user_" + "".join(ch if ch.isalnum() or ch in "._" else "_" for ch in raw).lower()

        username = base
        i = 1
        while User.objects.filter(username=username).exists():
            i += 1
            username = f"{base}_{i}"
        return username
